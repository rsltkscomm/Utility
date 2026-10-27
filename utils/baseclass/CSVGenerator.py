import csv
import random
from pathlib import Path
from datetime import datetime, timedelta
from faker import Faker
from collections import OrderedDict, defaultdict
import re
import openpyxl

from utils.constants.framework_constants import FrameworkConstants
from utils.ini_file_reader.config_reader import ConfigReader

faker = Faker()

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("ascii", "replace").decode("ascii"))

# Shared variables (equivalent to static fields)
csv_audience_count = 0
excel_audience_count = 0
base_audience_list = []
company_names = []


class CSVGenerator:
    csv_audience_count = 0
    excel_audience_count = 0

    # -------------------------------
    # Generate Normal CSV
    # -------------------------------
    @staticmethod
    def generate_csv(count, headers, partial_fill_attribute, partial_fill_count, file_path):
        all_data = []

        hobbies_list = ["Reading", "Traveling", "Sports", "Music"]
        cities_list = ["New York", "London", "Paris", "Berlin"]

        with open(file_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(headers)

            for i in range(count):
                row = {}

                for header in headers:
                    key = header.strip().lower()
                    value = ""

                    if key == "name":
                        value = faker.name()
                    elif key in ["email", "email id", "emailid"]:
                        value = faker.email()
                    elif key == "mobileno":
                        if partial_fill_attribute.lower() == "mobileno" and i >= partial_fill_count:
                            value = ""
                        else:
                            value = "91" + faker.msisdn()[:10]
                    elif key == "city":
                        if partial_fill_attribute.lower() == "city" and i >= partial_fill_count:
                            value = ""
                        else:
                            value = faker.city()
                    elif key == "hobbies":
                        value = random.choice(hobbies_list)
                    elif key == "country":
                        value = random.choice(cities_list)
                    elif key == "designation":
                        value = faker.job()
                    elif key == "companytitle":
                        value = faker.company()
                        company_names.append(value)
                    elif key == "id":
                        value = str(random.randint(100000, 999999))
                    elif key == "gender":
                        value = random.choice(["Male", "Female", "Other"])

                    row[header] = value

                writer.writerow(row.values())
                all_data.append(row)

        print(f"✅ CSV generated at: {file_path}")
        return all_data

    # -------------------------------
    # Generate Duplicate CSV
    # -------------------------------
    @staticmethod
    def generate_duplicate_csv(count, headers, partial_fill_attribute, partial_fill_count, file_path):
        all_data = []
        special_chars = "!%^$#@"

        unique_count = count // 2 if count > 10 else count
        duplicate_count = count - unique_count if count > 10 else count

        # Step 1: Unique Data
        unique_data = CSVGenerator.generate_csv(
            unique_count, headers, partial_fill_attribute, partial_fill_count, file_path
        )

        all_data.extend(unique_data)

        # Step 2: Duplicate Data
        with open(file_path, "a", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)

            for _ in range(duplicate_count):
                row = random.choice(unique_data)
                writer.writerow(row.values())
                all_data.append(row.copy())

        print(f"Total Rows: {len(all_data)}")
        return all_data

    # -------------------------------
    # Fill Missing Values
    # -------------------------------
    @staticmethod
    def fill_missing_values(data, attribute, file_path):
        for row in data:
            value = row.get(attribute, "").strip()

            if not value:
                if attribute.lower() == "city":
                    row[attribute] = faker.city()
                elif attribute.lower() == "name":
                    row[attribute] = faker.name()
                elif attribute.lower() in ["email", "emailid"]:
                    row[attribute] = faker.email()
                elif attribute.lower() == "mobileno":
                    row[attribute] = "91" + faker.msisdn()[:10]
                else:
                    row[attribute] = faker.word()

        CSVGenerator.write_csv(data, list(data[0].keys()), file_path)
        return data

    # -------------------------------
    # Write CSV
    # -------------------------------
    @staticmethod
    def write_csv(data, headers, file_path):
        with open(file_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(headers)

            for row in data:
                writer.writerow([row.get(h, "") for h in headers])

        print(f"✅ CSV written: {file_path}")

    # -------------------------------
    # Generate Base Audience
    # -------------------------------
    @staticmethod
    def generate_base_audience(header_key, list_type):
        csv_audience_count = 0
        base_audience_list.clear()

        prop = {}
        file_path = FrameworkConstants.get_dynamic_file_path()

        with open(file_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()

                if not line or line.startswith("#"):
                    continue

                if "=" not in line:
                    continue

                k, v = line.split("=", 1)
                prop[k.strip()] = v.strip()

        if header_key not in prop:
            raise Exception(f"Header key '{header_key}' not found in properties")

        headers = [h.strip() for h in prop[header_key].split(",")]

        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        csv_file = Path(FrameworkConstants.DYNAMIC_PATH) / f"BaseAudience{list_type.replace(' ', '')}{timestamp}.csv"

        with open(csv_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(headers)

            row_keys = sorted(
                [k for k in prop if k.lower().startswith("row")],
                key=lambda x: int(re.sub(r"\D", "", x))
            )

            for key in row_keys:
                values = prop[key].split("|")

                row_data = []
                row_map = OrderedDict()

                for i, header in enumerate(headers):
                    value = values[i].strip() if i < len(values) else ""

                    if header.lower() == "mobileno" and value.isdigit():
                        value = "\t" + value

                    row_data.append(value)
                    row_map[header] = value

                writer.writerow(row_data)
                base_audience_list.append(row_map)
                csv_audience_count += 1

        CSVGenerator.csv_audience_count = csv_audience_count
        globals()["csv_audience_count"] = csv_audience_count
        print(f"✅ CSV created at: {csv_file}")
        return str(csv_file)

    # -------------------------------
    # Generate Base Audience Excel
    # -------------------------------
    @staticmethod
    def generate_base_audience_excel(header_key, list_type):
        excel_audience_count = 0
        base_audience_list.clear()

        prop = {}
        file_path = FrameworkConstants.get_dynamic_file_path()

        with open(file_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()

                if not line or line.startswith("#"):
                    continue

                if "=" not in line:
                    continue

                k, v = line.split("=", 1)
                prop[k.strip()] = v.strip()

        if header_key not in prop:
            raise Exception(f"Header key '{header_key}' not found in properties")

        headers = [h.strip() for h in prop[header_key].split(",")]

        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        excel_file = Path(FrameworkConstants.DYNAMIC_PATH) / f"BaseAudience{list_type.replace(' ', '')}{timestamp}.xlsx"
        excel_file.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(headers)

        row_keys = sorted(
            [k for k in prop if k.lower().startswith("row")],
            key=lambda x: int(re.sub(r"\D", "", x))
        )

        for key in row_keys:
            values = prop[key].split("|")

            row_data = []
            row_map = OrderedDict()

            for i, header in enumerate(headers):
                value = values[i].strip() if i < len(values) else ""

                row_data.append(value)
                row_map[header] = value

            ws.append(row_data)
            base_audience_list.append(row_map)
            excel_audience_count += 1

        wb.save(str(excel_file))
        CSVGenerator.excel_audience_count = excel_audience_count
        CSVGenerator.csv_audience_count = excel_audience_count
        globals()["excel_audience_count"] = excel_audience_count
        globals()["csv_audience_count"] = excel_audience_count
        _safe_print(f"✅ Excel created at: {excel_file} with {excel_audience_count} rows")
        return str(excel_file)

    # -------------------------------
    # Generate Bulk Audience CSV
    # -------------------------------
    @staticmethod
    def generate_bulk_audience(header_key, list_type, count, partial_fill_count=20):
        count = int(count)
        base_audience_list.clear()

        prop = {}
        file_path = FrameworkConstants.get_dynamic_file_path()

        with open(file_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                k, v = line.split("=", 1)
                prop[k.strip()] = v.strip()

        if header_key not in prop:
            raise Exception(f"Header key '{header_key}' not found in properties")

        headers = [h.strip() for h in prop[header_key].split(",")]

        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        csv_file = Path(FrameworkConstants.DYNAMIC_PATH) / f"BulkAudience{list_type.replace(' ', '')}{timestamp}.csv"
        csv_file.parent.mkdir(parents=True, exist_ok=True)

        row_keys = sorted(
            [k for k in prop if k.lower().startswith("row")],
            key=lambda x: int(re.sub(r"\D", "", x))
        )

        template_rows = []
        for key in row_keys:
            values = prop[key].split("|")
            template_rows.append(values)

        with open(csv_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(headers)

            for i in range(count):
                base_vals = template_rows[i % len(template_rows)] if template_rows else []
                row_data = []
                row_map = OrderedDict()

                for j, header in enumerate(headers):
                    key = header.strip().lower()
                    if key in ["email", "email id", "emailid"]:
                        value = f"user_{timestamp}_{i}_{faker.email()}"
                    elif key == "mobileno":
                        value = "\t91" + str(random.randint(6000000000, 9999999999))
                    elif key == "name":
                        value = faker.name()
                    else:
                        value = base_vals[j].strip() if j < len(base_vals) else faker.word()

                    row_data.append(value)
                    row_map[header] = value

                writer.writerow(row_data)
                base_audience_list.append(row_map)

        CSVGenerator.csv_audience_count = count
        globals()["csv_audience_count"] = count
        _safe_print(f"✅ Bulk CSV created at: {csv_file}")
        return str(csv_file)

    # -------------------------------
    # Write Excel
    # -------------------------------
    @staticmethod
    def write_excel(data, headers, file_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(headers)

        for row in data:
            ws.append([row.get(h, "") for h in headers])

        wb.save(str(file_path))
        _safe_print(f"✅ Excel written: {file_path}")

    # -------------------------------
    # Get Company Names
    # -------------------------------
    @staticmethod
    def get_company_names():
        return company_names

    @staticmethod
    def generate_grouped_csvs(group_by_columns, limit=0):

        if isinstance(group_by_columns, str):
            group_by_columns = [col.strip() for col in group_by_columns.split(",")]

        base_dir = Path(FrameworkConstants.DYNAMIC_PATH) / "grouped"
        data_file = Path(FrameworkConstants.DYNAMIC_PATH) / "audiencedata.ini"

        base_dir.mkdir(parents=True, exist_ok=True)

        for f in base_dir.glob("*"):
            f.unlink()

        base_order = ConfigReader.get_property("baseorder", "").split(",")
        base_order = [x.strip() for x in base_order]

        grouped_data = defaultdict(list)
        created_files = []

        prop = {}
        with open(data_file, "r", encoding="utf-8") as f:
            for line in f:
                if "=" in line:
                    k, v = line.split("=", 1)
                    prop[k.strip()] = v.strip()

        for key, value in prop.items():
            if not key.lower().startswith("row"):
                continue

            values = value.split("|")
            if len(values) < len(base_order):
                continue

            row = [v.strip() for v in values[:len(base_order)]]

            group_key_parts = []
            for col in group_by_columns:
                if col in base_order:
                    idx = base_order.index(col)
                    clean_val = "".join(c if c.isalnum() else "_" for c in row[idx])
                    group_key_parts.append(clean_val)

            group_key = "_".join(group_key_parts)
            grouped_data[group_key].append(row)

        import time

        created = 0
        for group, rows in grouped_data.items():

            if limit > 0 and created >= limit:
                break

            file_path = base_dir / f"Group_{group}_{time.time_ns()}.csv"

            with open(file_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(group_by_columns)

                for row in rows:
                    selected = []
                    for col in group_by_columns:
                        idx = base_order.index(col)
                        selected.append(row[idx])
                    writer.writerow(selected)

            print(f"✅ Created group CSV: {file_path}")
            created_files.append(str(file_path))
            created += 1

        return created_files

    # -------------------------------
    # Generate Grouped Excels
    # -------------------------------
    @staticmethod
    def generate_grouped_excels(group_by_columns, limit=0):

        if isinstance(group_by_columns, str):
            group_by_columns = [col.strip() for col in group_by_columns.split(",")]

        base_dir = Path(FrameworkConstants.DYNAMIC_PATH) / "grouped"
        data_file = Path(FrameworkConstants.DYNAMIC_PATH) / "audiencedata.ini"

        base_dir.mkdir(parents=True, exist_ok=True)

        for f in base_dir.glob("*.xlsx"):
            try:
                f.unlink()
            except Exception:
                pass

        base_order = ConfigReader.get_property("baseorder", "").split(",")
        base_order = [x.strip() for x in base_order]

        grouped_data = defaultdict(list)
        created_files = []

        prop = {}
        with open(data_file, "r", encoding="utf-8") as f:
            for line in f:
                if "=" in line:
                    k, v = line.split("=", 1)
                    prop[k.strip()] = v.strip()

        for key, value in prop.items():
            if not key.lower().startswith("row"):
                continue

            values = value.split("|")
            if len(values) < len(base_order):
                continue

            row = [v.strip() for v in values[:len(base_order)]]

            group_key_parts = []
            for col in group_by_columns:
                if col in base_order:
                    idx = base_order.index(col)
                    clean_val = "".join(c if c.isalnum() else "_" for c in row[idx])
                    group_key_parts.append(clean_val)

            group_key = "_".join(group_key_parts)
            grouped_data[group_key].append(row)

        import time

        created = 0
        for group, rows in grouped_data.items():

            if limit > 0 and created >= limit:
                break

            file_path = base_dir / f"Group_{group}_{time.time_ns()}.xlsx"

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws.append(group_by_columns)

            for row in rows:
                selected = []
                for col in group_by_columns:
                    idx = base_order.index(col)
                    selected.append(row[idx])
                ws.append(selected)

            wb.save(str(file_path))
            _safe_print(f"✅ Created group Excel: {file_path}")
            created_files.append(str(file_path))
            created += 1

        return created_files

    # -------------------------------
    # Row Match Logic
    # -------------------------------
    @staticmethod
    def rows_match(base_row, uploaded_row, csv_header, base_order):
        for col in csv_header:
            if col in base_order:
                idx = base_order.index(col)
                if base_row[idx] != uploaded_row[idx]:
                    return False
        return True

        # -------------------------------
        # Inclusion Match Count
        # -------------------------------

    @staticmethod
    def get_inclusion_match_count_from_csvs(csv_paths):
        data_file = Path(FrameworkConstants.DYNAMIC_PATH) / "audiencedata.ini"
        base_order = ConfigReader.get_property("baseorder", "").split(",")
        base_order = [x.strip() for x in base_order]

        base_rows = []
        prop = {}

        # Load base data
        with open(data_file, "r", encoding="utf-8") as f:
            for line in f:
                if "=" in line:
                    k, v = line.split("=", 1)
                    prop[k.strip()] = v.strip()

        for key, value in prop.items():
            if not key.lower().startswith("row"):
                continue

            values = value.split("|")
            if len(values) < len(base_order):
                continue

            row = [v.strip() for v in values[:len(base_order)]]
            base_rows.append(row)

        total_match_count = 0

        # Process each CSV
        for csv_path in csv_paths:
            with open(csv_path, "r", encoding="utf-8") as f:
                reader = csv.reader(f)
                header = next(reader, None)

                if not header:
                    continue

                for row_values in reader:
                    uploaded_row = [""] * len(base_order)

                    for i, col in enumerate(header):
                        if col in base_order:
                            idx = base_order.index(col)
                            uploaded_row[idx] = row_values[i].strip()

                    for base_row in base_rows:
                        if CSVGenerator.rows_match(base_row, uploaded_row, header, base_order):
                            total_match_count += 1
                            break

        return total_match_count

    # -------------------------------
    # Inclusion Match Count From Excels
    # -------------------------------
    @staticmethod
    def get_inclusion_match_count_from_excels(excel_paths):
        data_file = Path(FrameworkConstants.DYNAMIC_PATH) / "audiencedata.ini"
        base_order = ConfigReader.get_property("baseorder", "").split(",")
        base_order = [x.strip() for x in base_order]

        base_rows = []
        prop = {}

        # Load base data
        with open(data_file, "r", encoding="utf-8") as f:
            for line in f:
                if "=" in line:
                    k, v = line.split("=", 1)
                    prop[k.strip()] = v.strip()

        for key, value in prop.items():
            if not key.lower().startswith("row"):
                continue

            values = value.split("|")
            if len(values) < len(base_order):
                continue

            row = [v.strip() for v in values[:len(base_order)]]
            base_rows.append(row)

        total_match_count = 0

        # Process each Excel file
        for excel_path in excel_paths:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            header = next(rows_iter, None)

            if not header:
                continue

            header = [str(c).strip() if c is not None else "" for c in header]

            for row_values in rows_iter:
                if not any(row_values):
                    continue
                uploaded_row = [""] * len(base_order)

                for i, col in enumerate(header):
                    if col in base_order and i < len(row_values):
                        idx = base_order.index(col)
                        val = row_values[i]
                        uploaded_row[idx] = str(val).strip() if val is not None else ""

                for base_row in base_rows:
                    if CSVGenerator.rows_match(base_row, uploaded_row, header, base_order):
                        total_match_count += 1
                        break

        return total_match_count
